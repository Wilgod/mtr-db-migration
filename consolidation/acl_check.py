"""
Folder ACL check: the NTFS permissions (Get-Acl, as permission.ps1 reads them)
of a folder and, optionally, its subfolders and files, read through Windows PowerShell.

    proc = start(r"\\\\nas\\share\\folder", levels=-1, own_only=False, files=True)
    with XlsxReport(path, root) as report:
        for line in proc.stdout:              # one JSON object per line
            item = parse_line(line)           # {"k": "item" | "error" | "skip" | "done", ...}
            report.write(item)

List mode, start(..., paths=[...]), reads only the listed folders / files.
Backups: backup_start(find_backups(folder)) reads permission.ps1's .acl.bak files instead
of the folders (same output lines), so the ACL as it was before permission.ps1 can be reported.
Paths longer than 260 characters are read through the \\\\?\\ prefix. Nothing is
changed on disk: Get-Acl only reads.

Apply: a report (possibly edited in Excel) put back on the folders. Every path in
it gets exactly the explicit entries listed for it; entries not in the report are
removed, inheritance is broken / restored to match "Inheritance broken":

    report_root, items, problems = read_report(xlsx)
    proc = apply_start(items, dry_run=True)   # same output lines; dry run changes nothing
"""
import base64
import csv
import json
import os
import subprocess
import tempfile
from datetime import datetime

import xlsxwriter

NO_WINDOW = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0

# permission.ps1's protected accounts and the right each should have
PROTECTED_ACCOUNTS = {
    "NT AUTHORITY\\SYSTEM": "FullControl",
    "BUILTIN\\Administrators": "FullControl",
    "corp.mtrc.com\\Domain Admins": "FullControl",
    "CREATOR OWNER": "FullControl",
    "MTRCORP\\dlpchan": "FullControl",
    "MTRCORP\\leslieip": "FullControl",
    "MTRCORP\\kwanwils": "FullControl",
    "MTRCORP\\CWPSP_READ": "Read",
    "MTRCORP\\CWPSP_RW": "ReadAndExecute",
    "MTRCORP\\CWP_MIG_PRD": "FullControl",
    "MTRCORP\\CWP_ODC": "Read",
    "MTRCORP\\cwppsup": "ReadAndExecute",
    "MTRCORP\\CWP_DEV": "Read",
}

# Windows PowerShell 5.1 (Get-Acl); one JSON object per output line.
_SCRIPT = r"""
$ErrorActionPreference = 'Continue'; $ProgressPreference = 'SilentlyContinue'
[Console]::OutputEncoding = New-Object System.Text.UTF8Encoding $false
$levels = [int]$env:ACL_LEVELS
$ownOnly = $env:ACL_OWN_ONLY -eq '1'; $files = $env:ACL_FILES -eq '1'
# A trailing \ is dropped, except on a drive root such as C:\
function Trim($p) { $p = $p.Trim().Trim('"'); if ($p.Length -gt 3) { $p.TrimEnd('\') } else { $p } }
$root = Trim $env:ACL_ROOT
# \\?\ (\\?\UNC\ for \\server\share) lifts the 260-character path limit.
function Long($p) { if ($p.StartsWith('\\?\')) { $p } elseif ($p.StartsWith('\\')) { '\\?\UNC\' + $p.Substring(2) } else { '\\?\' + $p } }
function Short($p) { if ($p.StartsWith('\\?\UNC\')) { '\\' + $p.Substring(8) } elseif ($p.StartsWith('\\?\')) { $p.Substring(4) } else { $p } }
function Clean($m) { $m.Replace('\\?\UNC\', '\\').Replace('\\?\', '') }
function Emit($o) { [Console]::Out.WriteLine(($o | ConvertTo-Json -Compress -Depth 4)); [Console]::Out.Flush() }
function Read-One($p, $isFile) {
    try {
        $acl = Get-Acl -LiteralPath $p -ErrorAction Stop
        $aces = @($acl.Access | ForEach-Object {
            [ordered]@{ a = $_.IdentityReference.Value; r = $_.FileSystemRights.ToString()
                        t = $_.AccessControlType.ToString(); i = $_.IsInherited
                        f = $_.InheritanceFlags.ToString(); p = $_.PropagationFlags.ToString() } })
        $own = @($aces | Where-Object { -not $_.i }).Count
        if ($ownOnly -and $p -ne $longRoot -and -not $acl.AreAccessRulesProtected -and $own -eq 0) {
            Emit @{ k = 'skip'; file = $isFile }; return
        }
        Emit ([ordered]@{ k = 'item'; file = $isFile; path = (Short $p); owner = $acl.Owner
                          broken = $acl.AreAccessRulesProtected; aces = $aces })
    } catch { Emit ([ordered]@{ k = 'error'; file = $isFile; path = (Short $p); error = (Clean $_.Exception.Message) }) }
}
if ($env:ACL_LIST) {
    # List mode: exactly the listed folders / files (one per line, UTF-8 file), nothing below them.
    $list = [System.IO.File]::ReadAllLines($env:ACL_LIST, [System.Text.Encoding]::UTF8)
    Remove-Item -LiteralPath $env:ACL_LIST -Force -ErrorAction SilentlyContinue
    $ownOnly = $false
    foreach ($line in $list) {
        $p = Trim $line
        if (-not $p) { continue }
        $lp = Long $p
        if ([System.IO.Directory]::Exists($lp)) { Read-One $lp $false }
        elseif ([System.IO.File]::Exists($lp)) { Read-One $lp $true }
        else { Emit ([ordered]@{ k = 'error'; file = $false; path = $p; error = 'Not found, or no access to it.' }) }
    }
    Emit @{ k = 'done' }; exit
}
$longRoot = Long $root
if (-not [System.IO.Directory]::Exists($longRoot)) {
    Emit ([ordered]@{ k = 'error'; file = $false; path = $root; error = 'Folder not found, or no access to it.' }); exit
}
Read-One $longRoot $false
if ($levels -ne 0) {
    $gci = @{ LiteralPath = $longRoot; Recurse = $true; Force = $true
              ErrorAction = 'SilentlyContinue'; ErrorVariable = 'listErrors' }
    if (-not $files) { $gci.Directory = $true }
    if ($levels -gt 0) { $gci.Depth = $levels - 1 }
    Get-ChildItem @gci | ForEach-Object { Read-One $_.FullName (-not $_.PSIsContainer) }
    foreach ($e in $listErrors) {
        Emit ([ordered]@{ k = 'error'; file = $false; path = (Short "$($e.TargetObject)"); error = (Clean $e.Exception.Message) })
    }
} elseif ($files) {
    Get-ChildItem -LiteralPath $longRoot -File -Force -ErrorAction SilentlyContinue | ForEach-Object { Read-One $_.FullName $true }
}
Emit @{ k = 'done' }
"""

# Generic rights (CREATOR OWNER, inherit-only entries) that ToString() leaves as a number
_GENERIC = {"268435456": "FullControl (generic)", "-536805376": "Modify (generic)",
            "-1610612736": "ReadAndExecute (generic)", "1073741824": "Write (generic)",
            "-2147483648": "Read (generic)"}


def start(root, levels, own_only, files, paths=None):
    """Windows PowerShell reading the ACLs; levels: 0 = this folder only (plus its files), -1 = all.
With `paths` (list mode) only those folders / files are read and the other arguments are ignored."""
    # PowerShell 7's module path breaks Windows PowerShell's Get-Acl (Microsoft.PowerShell.Security)
    env = {k: v for k, v in os.environ.items() if k.upper() != "PSMODULEPATH"}
    env.update(ACL_ROOT=root, ACL_LEVELS=str(levels), ACL_OWN_ONLY="1" if own_only else "0",
               ACL_FILES="1" if files else "0")
    if paths:
        # the list goes in a file: a long list does not fit in an environment variable
        fd, list_file = tempfile.mkstemp(prefix="acl_paths_", suffix=".txt")
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write("\n".join(paths))
        env["ACL_LIST"] = list_file
    encoded = base64.b64encode(_SCRIPT.encode("utf-16-le")).decode("ascii")
    return subprocess.Popen(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                             "-EncodedCommand", encoded], env=env, creationflags=NO_WINDOW,
                            stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)


# Windows PowerShell 5.1: permission.ps1's ACL backups ($acl | ConvertTo-Csv, one <path>.acl.bak
# per folder) -> the same item lines as _SCRIPT. The Sddl column holds the whole ACL; accounts are
# named now (SID -> name), or from the backup's AccessToString when the SID no longer resolves.
_BACKUP_SCRIPT = r"""
$ErrorActionPreference = 'Continue'; $ProgressPreference = 'SilentlyContinue'
[Console]::OutputEncoding = New-Object System.Text.UTF8Encoding $false
$list = [System.IO.File]::ReadAllLines($env:ACL_LIST, [System.Text.Encoding]::UTF8)
Remove-Item -LiteralPath $env:ACL_LIST -Force -ErrorAction SilentlyContinue
function Long($p) { if ($p.StartsWith('\\?\')) { $p } elseif ($p.StartsWith('\\')) { '\\?\UNC\' + $p.Substring(2) } else { '\\?\' + $p } }
function Emit($o) { [Console]::Out.WriteLine(($o | ConvertTo-Json -Compress -Depth 4)); [Console]::Out.Flush() }
$sidType = [System.Security.Principal.SecurityIdentifier]
$rids = @{ DA = 512; DU = 513; DG = 514; DC = 515; DD = 516; CA = 517; SA = 518; EA = 519; PA = 520
           CN = 522; AP = 525; KA = 526; EK = 527; RS = 553; RO = 498 }
$names = @{}
function Name($sid) {
    if (-not $names.ContainsKey($sid.Value)) {
        try { $names[$sid.Value] = $sid.Translate([System.Security.Principal.NTAccount]).Value } catch { $names[$sid.Value] = $null }
    }
    $names[$sid.Value]
}
foreach ($f in $list) {
    if (-not $f.Trim()) { continue }
    $path = [System.IO.Path]::GetFileName($f)
    try {
        $row = @(Import-Csv -LiteralPath $f -Encoding UTF8)[0]
        if (-not $row -or -not $row.Sddl) { throw 'No Sddl column: not an ACL backup of permission.ps1 (Get-Acl | ConvertTo-Csv).' }
        if ($row.Path) { $path = $row.Path -replace '^.*::', '' }
        # Windows PowerShell writes "#TYPE System.Security.AccessControl.DirectorySecurity / FileSecurity";
        # PowerShell 7 does not: then the path on disk decides, else it is taken for a folder.
        $type = [System.IO.File]::ReadLines($f) | Select-Object -First 1
        $isFile = if ($type -match '^#TYPE') { $type -match 'FileSecurity' }
                  else { -not [System.IO.Directory]::Exists((Long $path)) -and [System.IO.File]::Exists((Long $path)) }
        $sec = if ($isFile) { New-Object System.Security.AccessControl.FileSecurity } else { New-Object System.Security.AccessControl.DirectorySecurity }
        # Domain aliases (DA = Domain Admins, DU = Domain Users...) only parse on a computer of that
        # domain: written out as full SIDs, with the domain SID found in the same SDDL.
        $sddl = $row.Sddl
        $dom = [regex]::Match($sddl, 'S-1-5-21-\d+-\d+-\d+').Value
        if ($dom) {
            $sddl = [regex]::Replace($sddl, '(?<=;|O:|G:)(DA|DU|DG|DC|DD|CA|SA|EA|PA|CN|AP|KA|EK|RS|RO)(?=\)|[GDS]:|$)',
                                     { param($m) "$dom-$($rids[$m.Value])" })
        }
        $sec.SetSecurityDescriptorSddlForm($sddl)
        $rules = @($sec.GetAccessRules($true, $true, $sidType))
        # AccessToString: one "<account> Allow|Deny  <rights>" line per rule, in the same order
        $saved = @("$($row.AccessToString)" -split '\r?\n' | Where-Object { $_ } | ForEach-Object {
            if ($_ -match '^(.+?) (Allow|Deny)  ') { $Matches[1] } else { $null } })
        $i = 0
        $aces = @($rules | ForEach-Object {
            $a = Name $_.IdentityReference
            if (-not $a) { $a = if ($saved.Count -eq $rules.Count -and $saved[$i]) { $saved[$i] } else { $_.IdentityReference.Value } }
            $i++
            [ordered]@{ a = $a; r = $_.FileSystemRights.ToString(); t = $_.AccessControlType.ToString(); i = $_.IsInherited
                        f = $_.InheritanceFlags.ToString(); p = $_.PropagationFlags.ToString() } })
        $owner = if ($row.Owner) { $row.Owner } else { $o = $sec.GetOwner($sidType); if ($o) { $n = Name $o; if ($n) { $n } else { $o.Value } } else { '' } }
        Emit ([ordered]@{ k = 'item'; file = [bool]$isFile; path = $path; owner = $owner
                          broken = $sec.AreAccessRulesProtected; aces = $aces })
    } catch { Emit ([ordered]@{ k = 'error'; file = $false; path = $path; error = "$f : $($_.Exception.Message)" }) }
}
Emit @{ k = 'done' }
"""


def find_backups(folder):
    """permission.ps1's ACL backup files (*.acl.bak) in `folder` and below, sorted (parents first)."""
    found = []
    for dirpath, _, names in os.walk(folder):
        found += [os.path.join(dirpath, n) for n in names if n.lower().endswith(".acl.bak")]
    return sorted(found, key=str.lower)


def backup_start(files):
    """Windows PowerShell reading permission.ps1's .acl.bak files; output lines as start()."""
    env = {k: v for k, v in os.environ.items() if k.upper() != "PSMODULEPATH"}
    fd, list_file = tempfile.mkstemp(prefix="acl_backups_", suffix=".txt")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write("\n".join(files))
    env["ACL_LIST"] = list_file
    encoded = base64.b64encode(_BACKUP_SCRIPT.encode("utf-16-le")).decode("ascii")
    return subprocess.Popen(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                             "-EncodedCommand", encoded], env=env, creationflags=NO_WINDOW,
                            stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)


def clean_paths(lines):
    """Pasted lines -> paths: trimmed, quotes removed, blanks and repeats dropped, order kept."""
    out, seen = [], set()
    for line in lines:
        p = line.strip().strip('"').strip()
        if p and p.lower() not in seen:
            seen.add(p.lower())
            out.append(p)
    return out


def _looks_like_path(text):
    return text.startswith("\\\\") or (len(text) > 2 and text[1] == ":" and text[2] in "\\/")


def read_path_list(path):
    """Paths from a .txt (one per line), .csv or .xlsx (any cell of any sheet): every value that
looks like a path (\\\\server\\share\\... or X:\\...); headers and other columns are ignored."""
    values = []
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values += [str(v).strip() for v in row if isinstance(v, str)]
        finally:
            book.close()
    elif path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
            for row in csv.reader(f):
                values += row
    else:
        with open(path, encoding="utf-8-sig", errors="replace") as f:
            values = f.read().splitlines()
    return clean_paths(v for v in values if _looks_like_path(v.strip().strip('"')))


def parse_line(raw):
    """One output line -> dict, or None for anything that is not ours (PowerShell noise)."""
    line = raw.decode("utf-8", errors="replace").strip()
    if not line.startswith("{"):
        return None
    try:
        item = json.loads(line)
    except ValueError:
        return None
    if item.get("k") == "item":
        item["aces"] = item.get("aces") or []   # ConvertTo-Json writes an empty array as null
    return item


def rights(ace):
    return _GENERIC.get(ace["r"], ace["r"])


def applies_to(ace, is_file=False):
    """InheritanceFlags / PropagationFlags as the Windows security dialog words them."""
    if is_file:
        return "This file"
    f, p = ace["f"], ace["p"]
    ci, oi = "ContainerInherit" in f, "ObjectInherit" in f
    if "InheritOnly" in p:
        text = {(True, True): "Subfolders and files only", (True, False): "Subfolders only",
                (False, True): "Files only"}.get((ci, oi), "(nothing)")
    else:
        text = {(True, True): "This folder, subfolders and files", (True, False): "This folder and subfolders",
                (False, True): "This folder and files"}.get((ci, oi), "This folder only")
    return text + (" (one level down)" if "NoPropagateInherit" in p else "")


def _short(account):
    return account.split("\\")[-1].lower()


def protected(account):
    """Is `account` one of PROTECTED_ACCOUNTS (the domain part may differ: MTRCORP vs corp.mtrc.com)."""
    return any(_short(account) == _short(p) for p in PROTECTED_ACCOUNTS)


def missing_protected(item):
    """PROTECTED_ACCOUNTS with no Allow entry on `item` (a folder or file)."""
    have = {_short(a["a"]) for a in item["aces"] if a["t"] == "Allow"}
    return [p for p in PROTECTED_ACCOUNTS if _short(p) not in have]


class XlsxReport:
    """Excel report, one row per entry, streamed as items arrive (constant memory, so a whole
NAS share fits): sheet "Summary", then "ACL", "ACL 2"... of at most SHEET_ROWS rows each."""
    HEADER = ["Path", "Item type", "Owner", "Inheritance broken", "Account", "Rights", "Allow / Deny",
              "Inherited", "Applies to", "Protected account", "Missing protected accounts", "Error"]
    WIDTHS = [90, 9, 24, 10, 36, 34, 11, 9, 30, 10, 50, 60]
    SHEET_ROWS = 1_000_000              # Excel's limit is 1,048,576

    def __init__(self, path, root):
        with open(path, "wb"):          # fail now (file open in Excel?), not after the whole read
            pass
        self.root, self.started = root, datetime.now()
        self.book = xlsxwriter.Workbook(path, {"constant_memory": True, "strings_to_formulas": False,
                                               "strings_to_urls": False, "strings_to_numbers": False})
        self.head = self.book.add_format({"bold": True, "bg_color": "#DDEBF7", "border": 1})
        self.red = self.book.add_format({"font_color": "#C00000"})
        self.summary = self.book.add_worksheet("Summary")
        self.counts = {"Folders": 0, "Files": 0, "Errors": 0, "Missing a protected account": 0, "Entries": 0}
        self.sheet, self.row, self.sheets = None, 0, 0

    def _row(self, values, fmt=None):
        if self.sheet is None or self.row > self.SHEET_ROWS:
            self._finish_sheet()
            self.sheets += 1
            self.sheet = self.book.add_worksheet("ACL" if self.sheets == 1 else f"ACL {self.sheets}")
            for i, w in enumerate(self.WIDTHS):
                self.sheet.set_column(i, i, w)
            self.sheet.write_row(0, 0, self.HEADER, self.head)
            self.sheet.freeze_panes(1, 0)
            self.row = 1
        for i, v in enumerate(values):
            if v != "":
                self.sheet.write_string(self.row, i, str(v), fmt)
        self.row += 1

    def _finish_sheet(self):
        if self.sheet is not None:
            self.sheet.autofilter(0, 0, self.row - 1, len(self.HEADER) - 1)

    def write(self, item):
        kind = "File" if item.get("file") else "Folder"
        if item["k"] == "error":
            self.counts["Errors"] += 1
            self._row([item["path"], kind] + [""] * 9 + [item["error"]], self.red)
            return
        if item["k"] != "item":
            return
        self.counts["Files" if item.get("file") else "Folders"] += 1
        missing = "; ".join(missing_protected(item))
        self.counts["Missing a protected account"] += bool(missing)
        broken = "yes" if item["broken"] else "no"
        if not item["aces"]:
            self._row([item["path"], kind, item["owner"], broken] + [""] * 6 + [missing, ""])
        for a in item["aces"]:
            self.counts["Entries"] += 1
            self._row([item["path"], kind, item["owner"], broken, a["a"], rights(a), a["t"],
                       "yes" if a["i"] else "no", applies_to(a, item.get("file")),
                       "yes" if protected(a["a"]) else "", missing, ""],
                      self.red if a["t"] == "Deny" else None)

    def close(self):
        self._finish_sheet()
        if self.sheet is None:          # nothing read: still an (empty) ACL sheet
            self._row([])
            self._finish_sheet()
        s = self.summary
        s.set_column(0, 0, 30)
        s.set_column(1, 1, 90)
        rows = [("Checked", self.root), ("Started", f"{self.started:%Y-%m-%d %H:%M:%S}"),
                ("Finished", f"{datetime.now():%Y-%m-%d %H:%M:%S}")] + [(k, f"{v:,}") for k, v in self.counts.items()]
        rows.append(("Protected accounts (permission.ps1)", ", ".join(PROTECTED_ACCOUNTS)))
        for r, (k, v) in enumerate(rows):
            s.write_string(r, 0, k, self.head)
            s.write_string(r, 1, v)
        self.book.close()

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.close()


# ---- Apply a report back to the folders ------------------------------------------------------

# FileSystemRights names (as FileSystemRights.ToString() writes them) -> access mask
_RIGHT_BITS = {
    "readdata": 0x1, "listdirectory": 0x1, "writedata": 0x2, "createfiles": 0x2, "appenddata": 0x4,
    "createdirectories": 0x4, "readextendedattributes": 0x8, "writeextendedattributes": 0x10,
    "executefile": 0x20, "traverse": 0x20, "deletesubdirectoriesandfiles": 0x40, "readattributes": 0x80,
    "writeattributes": 0x100, "write": 0x116, "delete": 0x10000, "readpermissions": 0x20000,
    "read": 0x20089, "readandexecute": 0x200A9, "modify": 0x301BF, "changepermissions": 0x40000,
    "takeownership": 0x80000, "synchronize": 0x100000, "fullcontrol": 0x1F01FF,
}
# GENERIC_READ / WRITE / EXECUTE / ALL -> the file rights they stand for (a rule cannot hold generic bits)
_GENERIC_BITS = ((0x80000000, 0x120089), (0x40000000, 0x120116), (0x20000000, 0x1200A0), (0x10000000, 0x1F01FF))
_APPLIES = {
    "this folder only": ("None", "None"),
    "this folder, subfolders and files": ("ContainerInherit, ObjectInherit", "None"),
    "this folder and subfolders": ("ContainerInherit", "None"),
    "this folder and files": ("ObjectInherit", "None"),
    "subfolders and files only": ("ContainerInherit, ObjectInherit", "InheritOnly"),
    "subfolders only": ("ContainerInherit", "InheritOnly"),
    "files only": ("ObjectInherit", "InheritOnly"),
    "this file": ("None", "None"),
}


def parse_rights(text):
    """Rights cell (e.g. "Modify, Synchronize", "FullControl (generic)", "-1610612736") -> access mask."""
    t = str(text).strip()
    generic = {v: k for k, v in _GENERIC.items()}
    t = generic.get(t, t)
    if t.lstrip("-").isdigit():
        raw, mask = int(t) & 0xFFFFFFFF, int(t) & 0x0FFFFFFF
        for bit, rights_ in _GENERIC_BITS:
            if raw & bit:
                mask |= rights_
    else:
        mask = 0
        for name in t.split(","):
            name = name.strip().lower()
            if name not in _RIGHT_BITS:
                raise ValueError(f"unknown Rights '{text}' (e.g. FullControl, Modify, ReadAndExecute, Read, Write)")
            mask |= _RIGHT_BITS[name]
    if not mask & 0x1F01FF:
        raise ValueError(f"Rights '{text}' grant nothing")
    return mask & 0x1F01FF


def parse_applies_to(text):
    """"Applies to" cell (applies_to()'s wording) -> (InheritanceFlags, PropagationFlags) as text."""
    t = str(text).strip().lower()
    one_level = t.endswith("(one level down)")
    t = t.replace("(one level down)", "").strip() or "this folder, subfolders and files"
    if t not in _APPLIES:
        raise ValueError(f"unknown Applies to '{text}'")
    f, p = _APPLIES[t]
    if one_level and f != "None":
        p = "NoPropagateInherit" if p == "None" else p + ", NoPropagateInherit"
    return f, p


def read_report(path):
    """An ACL report (.xlsx from 'Full report to Excel' / 'Export shown to Excel', edited or not) ->
(root it was taken from, [item], [problem]). item = {"path", "protect": True / False / None (keep),
"aces": [{"a", "t", "r" (mask), "f", "p", "i", "row"}]}, in report order (parents first).
Paths with a bad row are left out and listed in problems; so are rows with only a read error."""
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        root = ""
        if "Summary" in book.sheetnames:
            for row in book["Summary"].iter_rows(values_only=True):
                if len(row) > 1 and row[0] == "Checked":
                    root = str(row[1] or "").strip()
        sheets = [s for s in book.worksheets if s.title == "ACL" or s.title.startswith("ACL ")]
        if not sheets:
            raise ValueError("No 'ACL' sheet: choose a report made by 'Full report to Excel' or 'Export shown to Excel'.")
        items, bad = {}, {}
        for sheet in sheets:
            rows = sheet.iter_rows(values_only=True)
            head = [str(v or "").strip() for v in next(rows, ())]
            col = {h: i for i, h in enumerate(head) if h}
            missing = [h for h in ("Path", "Account", "Rights") if h not in col]
            if missing:
                raise ValueError(f"Sheet '{sheet.title}' has no column {', '.join(missing)}.")
            for n, row in enumerate(rows, start=2):
                def cell(h):
                    v = row[col[h]] if h in col and col[h] < len(row) else None
                    return "" if v is None else str(v).strip()
                p = cell("Path")
                if not p or (cell("Error") and not cell("Account")):
                    continue
                if len(p) > 3:
                    p = p.rstrip("\\")
                where = f"{sheet.title} row {n}"
                item = items.setdefault(p.lower(), {"path": p, "protect": None, "aces": []})
                try:
                    broken = cell("Inheritance broken").lower()
                    if broken:
                        if broken not in ("yes", "no"):
                            raise ValueError(f"Inheritance broken must be yes or no, not '{broken}'")
                        if item["protect"] is not None and item["protect"] != (broken == "yes"):
                            raise ValueError("Inheritance broken differs from this path's other rows")
                        item["protect"] = broken == "yes"
                    if not cell("Account"):
                        continue
                    t = cell("Allow / Deny").capitalize() or "Allow"
                    if t not in ("Allow", "Deny"):
                        raise ValueError(f"Allow / Deny must be Allow or Deny, not '{t}'")
                    f, pf = parse_applies_to(cell("Applies to"))
                    item["aces"].append({"a": cell("Account"), "t": t, "r": parse_rights(cell("Rights")),
                                         "f": f, "p": pf, "i": cell("Inherited").lower() == "yes", "row": where})
                except ValueError as e:
                    bad.setdefault(p.lower(), f"{p}: {where}: {e}")
        return root, [v for k, v in items.items() if k not in bad], list(bad.values())
    finally:
        book.close()


def rebase(items, old_root, new_root):
    """Items of a report taken from old_root, moved under new_root (e.g. the migrated copy) -> (items, problems)."""
    old = old_root.rstrip("\\") if len(old_root) > 3 else old_root
    new = new_root.rstrip("\\") if len(new_root) > 3 else new_root
    moved, problems = [], []
    for it in items:
        p = it["path"]
        if p.lower() == old.lower():
            moved.append(dict(it, path=new))
        elif p.lower().startswith(old.rstrip("\\").lower() + "\\"):
            moved.append(dict(it, path=new.rstrip("\\") + p[len(old.rstrip("\\")):]))
        else:
            problems.append(f"{p}: not under the report's folder {old_root}")
    return moved, problems


# Windows PowerShell 5.1: replaces each item's explicit entries (and inheritance) with the plan's.
# Output: one JSON object per line, {"k": "item" | "error" | "done", ...}.
_APPLY_SCRIPT = r"""
$ErrorActionPreference = 'Continue'; $ProgressPreference = 'SilentlyContinue'
[Console]::OutputEncoding = New-Object System.Text.UTF8Encoding $false
$dry = $env:ACL_DRY -eq '1'
# assigned first: Windows PowerShell's ConvertFrom-Json emits a JSON array as one object
$plan = [System.IO.File]::ReadAllText($env:ACL_PLAN, [System.Text.Encoding]::UTF8) | ConvertFrom-Json
Remove-Item -LiteralPath $env:ACL_PLAN -Force -ErrorAction SilentlyContinue
function Long($p) { if ($p.StartsWith('\\?\')) { $p } elseif ($p.StartsWith('\\')) { '\\?\UNC\' + $p.Substring(2) } else { '\\?\' + $p } }
function Clean($m) { $m.Replace('\\?\UNC\', '\\').Replace('\\?\', '') }
function Emit($o) { [Console]::Out.WriteLine(($o | ConvertTo-Json -Compress -Depth 5)); [Console]::Out.Flush() }
function Msg($e) { $x = $e.Exception; if ($x.InnerException) { $x = $x.InnerException }; Clean $x.Message }
$sidType = [System.Security.Principal.SecurityIdentifier]
$sids = @{}
function Sid($name) {
    if (-not $sids.ContainsKey($name)) {
        try {
            $sids[$name] = if ($name -match '^S-1-\d') { New-Object System.Security.Principal.SecurityIdentifier($name) }
                           else { ([System.Security.Principal.NTAccount]$name).Translate($sidType) }
        } catch { $sids[$name] = "Cannot resolve account '$name' (not found in AD?)" }
    }
    if ($sids[$name] -is [string]) { throw $sids[$name] }
    $sids[$name]
}
function Name($sid) { try { $sid.Translate([System.Security.Principal.NTAccount]).Value } catch { $sid.Value } }
# Generic bits mapped to file rights, Synchronize ignored: how the security dialog compares rights
function Mask($rights) {
    $u = [int64][int]$rights -band 0xFFFFFFFFL; $m = $u -band 0x0FFFFFFF
    if ($u -band 0x80000000L) { $m = $m -bor 0x120089 }
    if ($u -band 0x40000000L) { $m = $m -bor 0x120116 }
    if ($u -band 0x20000000L) { $m = $m -bor 0x1200A0 }
    if ($u -band 0x10000000L) { $m = $m -bor 0x1F01FF }
    $m -band 0x0F01FF
}
function Key($r, $isDir) {
    $f = if ($isDir) { [int]$r.InheritanceFlags } else { 0 }; $p = if ($isDir) { [int]$r.PropagationFlags } else { 0 }
    "$($r.IdentityReference.Value)|$($r.AccessControlType)|$(Mask $r.FileSystemRights)|$f|$p"
}
function Change($op, $r) {
    [ordered]@{ op = $op; a = (Name $r.IdentityReference); r = $r.FileSystemRights.ToString(); t = $r.AccessControlType.ToString()
                f = $r.InheritanceFlags.ToString(); p = $r.PropagationFlags.ToString() }
}
foreach ($it in $plan) {
    $lp = Long $it.path
    try {
        if ([System.IO.Directory]::Exists($lp)) { $fsi = New-Object System.IO.DirectoryInfo($lp); $isDir = $true }
        elseif ([System.IO.File]::Exists($lp)) { $fsi = New-Object System.IO.FileInfo($lp); $isDir = $false }
        else { Emit ([ordered]@{ k = 'error'; file = $false; path = $it.path; error = 'Not found, or no access to it.' }); continue }
        $sec = $fsi.GetAccessControl('Access')
        $before = $sec.GetSecurityDescriptorSddlForm('Access')
        $wasProt = $sec.AreAccessRulesProtected
        $prot = if ($null -eq $it.protect) { $wasProt } else { [bool]$it.protect }
        # While the item inherits, rows marked Inherited come from the parent and are not written;
        # with inheritance broken every row becomes an explicit entry.
        $want = if ($isDir) { New-Object System.Security.AccessControl.DirectorySecurity } else { New-Object System.Security.AccessControl.FileSecurity }
        foreach ($a in $it.aces) {
            if ($a.i -and -not $prot) { continue }
            $f = if ($isDir) { $a.f } else { 'None' }; $p = if ($isDir) { $a.p } else { 'None' }
            try {
                $want.AddAccessRule((New-Object System.Security.AccessControl.FileSystemAccessRule(
                    (Sid $a.a), [System.Security.AccessControl.FileSystemRights][int]$a.r,
                    [System.Security.AccessControl.InheritanceFlags]$f, [System.Security.AccessControl.PropagationFlags]$p,
                    [System.Security.AccessControl.AccessControlType]$a.t)))
            } catch { throw "$($a.row), $($a.a): $(Msg $_)" }
        }
        $have = @{}; foreach ($r in $sec.GetAccessRules($true, $false, $sidType)) { $have[(Key $r $isDir)] = $r }
        $need = @{}; foreach ($r in $want.GetAccessRules($true, $false, $sidType)) { $need[(Key $r $isDir)] = $r }
        $changes = @()
        if ($prot -ne $wasProt) { $changes += , [ordered]@{ op = $(if ($prot) { 'break' } else { 'inherit' }) } }
        foreach ($k in $need.Keys) { if (-not $have.ContainsKey($k)) { $changes += , (Change '+' $need[$k]) } }
        foreach ($k in $have.Keys) { if (-not $need.ContainsKey($k)) { $changes += , (Change '-' $have[$k]) } }
        if ($changes.Count -and -not $dry) {
            foreach ($r in @($sec.GetAccessRules($true, $false, $sidType))) { $sec.RemoveAccessRuleSpecific($r) }
            $sec.SetAccessRuleProtection($prot, $false)
            foreach ($r in $want.GetAccessRules($true, $false, $sidType)) { $sec.AddAccessRule($r) }
            $fsi.SetAccessControl($sec)
        }
        Emit ([ordered]@{ k = 'item'; file = (-not $isDir); path = $it.path; changes = $changes
                          before = $(if ($changes.Count) { $before } else { '' }) })
    } catch { Emit ([ordered]@{ k = 'error'; file = $false; path = $it.path; error = (Msg $_) }) }
}
Emit @{ k = 'done' }
"""


def apply_start(items, dry_run):
    """Windows PowerShell putting read_report() items on disk (dry_run: only compares). Output as start():
{"k": "item", "path", "file", "changes": [{"op": "+" | "-" | "break" | "inherit", ...}], "before": SDDL} per item."""
    env = {k: v for k, v in os.environ.items() if k.upper() != "PSMODULEPATH"}
    fd, plan_file = tempfile.mkstemp(prefix="acl_plan_", suffix=".json")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        json.dump([{k: it[k] for k in ("path", "protect", "aces")} for it in items], f)
    env.update(ACL_PLAN=plan_file, ACL_DRY="1" if dry_run else "0")
    encoded = base64.b64encode(_APPLY_SCRIPT.encode("utf-16-le")).decode("ascii")
    return subprocess.Popen(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                             "-EncodedCommand", encoded], env=env, creationflags=NO_WINDOW,
                            stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)


def change_text(c, is_file=False):
    """One change from apply_start() in words."""
    if c["op"] == "break":
        return "break inheritance (entries from the parent are no longer inherited)"
    if c["op"] == "inherit":
        return "restore inheritance from the parent"
    return f"{'add' if c['op'] == '+' else 'remove'} {c['a']} {c['t']} {rights(c)} ({applies_to(c, is_file)})"
