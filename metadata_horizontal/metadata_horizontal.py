"""
Horizontal metadata: one row per file for a fixed, config-driven set of metadata
keys, while every other key stays in the vertical "Metadata" table (one row per
key).

config.json (next to this file):

    "HorizontalMetadata": {
      "TableName": "MetadataHorizontal",
      "Columns": {                       # metadata key  ->  column name in the table
        "ishareFileType": "FileType",
        ...
      },
      "ColumnMaxLength": 4000,           # varchar length of every metadata column
      "CurrentVersionOnly": true,        # skip rows whose VersionLabel is not 'Current'
      "RemoveMovedKeysFromMetadata": true
    },
    "SkipColumns": [ ... ]               # source columns stored nowhere

Matching of horizontal keys / skip columns against file headers is case-insensitive.

Upload (CSV or XLSX) - import_file():
    every row  ->  matched to Files by FileId (only if the row's other identities
                   agree), else UniqueId (SharePointUniqueId), else
                   OriginalSharePointPath, else FileRef (path part of
                   OriginalSharePointPath). Folder rows (FSObjType=1) and old
                   versions are skipped; unmatched rows go to a report CSV.
    configured keys  ->  1 row in the horizontal table (upsert by FileId)
    other keys       ->  vertical "Metadata" rows (non-empty values only)
    For every matched file the old Metadata rows of all keys present in the file
    (horizontal + vertical) are replaced, so a re-upload never duplicates.
    One transaction per upload: it either fully applies or not at all.

Existing data - convert_existing():
    pivots configured keys already in "Metadata" into the horizontal table and
    (optionally) deletes them from "Metadata".

Matrix - build_matrix():
    rebuilds "MetadataMatrix" (config "MetadataMatrix.TableName"): one row per
    file, one text column per distinct "Metadata"."MetadataKey". "Metadata" is
    only read, never changed.

CLI:
    python metadata_horizontal.py ensure
    python metadata_horizontal.py upload  <file.csv|file.xlsx> [--sheet NAME]
    python metadata_horizontal.py convert [--keep-vertical]
    python metadata_horizontal.py ddl
    python metadata_horizontal.py matrix
  connection: --url postgresql://user:pw@host:5432/db  (or env PG_URL;
              default: config.json "Connection" + env PGPASSWORD;
              "Kerberos": true logs in as the Windows user, no password;
              "Role": "cwp_mig_prd" runs every session as SET ROLE cwp_mig_prd)
"""
import csv
import io
import json
import os
import re
import sys
import time
import urllib.parse
from datetime import datetime

from sqlalchemy import create_engine, text

import metadata_matrix as mm

BASE_COLUMNS = ("MetadataHorizontalId", "FileId", "SiteName", "SiteId",
                "SourceFile", "ImportedAt", "UpdatedAt")
ID_COLUMNS = ("FileId", "UniqueId", "FileRef", "OriginalSharePointPath",
              "FSObjType", "VersionLabel")
DEFAULT_TABLE = "MetadataHorizontal"
STAGE_BATCH = 5000          # rows per COPY into the staging table
_IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,62}$")
_ILLEGAL = re.compile(r"[\x00]")   # PostgreSQL text cannot hold NUL

if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BASE_DIR, "config.json")


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

def load_config(path=CONFIG_FILE):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_config(app_config, path=CONFIG_FILE):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(app_config, f, indent=2, ensure_ascii=False)


def role_name(raw):
    """Clean the "Role" setting: accepts "cwp_mig_prd" or a pasted "SET ROLE cwp_mig_prd;"."""
    role = (raw or "").strip().rstrip(";").strip()
    if role[:9].lower() == "set role ":
        role = role[9:].strip()
    if role and not all(ch.isalnum() or ch in "_$" for ch in role):
        raise ValueError("Role: enter just the role name, e.g. cwp_mig_prd (letters, digits, _ or $).")
    return role


def connection_url(app_config, password=None):
    c = app_config.get("Connection", {})
    user = urllib.parse.quote(c.get("User", "postgres"), safe="")
    pw = password if password is not None else os.environ.get("PGPASSWORD", "")
    if c.get("Kerberos"):
        pw = ""  # GSSAPI/SSPI ticket of the signed-in user; Host must be the server's full DNS name
    auth = user + (":" + urllib.parse.quote(pw, safe="") if pw else "")
    url = (f"postgresql+psycopg2://{auth}@{c.get('Host', 'localhost')}:{c.get('Port', 5432)}/"
           f"{urllib.parse.quote(c.get('Database', 'postgres'), safe='')}")
    params = {}
    if c.get("SslMode"):
        params["sslmode"] = c["SslMode"]
    if c.get("Kerberos"):
        params["krbsrvname"] = c.get("KrbSrvName") or "postgres"
    role = role_name(c.get("Role"))
    if role:
        params["options"] = f"-c role={role}"  # every session starts as SET ROLE <role>
    if params:
        url += "?" + urllib.parse.urlencode(params)
    return url


class HConfig:
    """Validated view of config.json for horizontal metadata."""

    def __init__(self, app_config):
        sec = app_config.get("HorizontalMetadata") or {}
        self.table = sec.get("TableName") or DEFAULT_TABLE
        self.columns = dict(sec.get("Columns") or {})          # key -> column
        self.max_len = int(sec.get("ColumnMaxLength") or 4000)
        self.current_only = bool(sec.get("CurrentVersionOnly", True))
        self.remove_moved = bool(sec.get("RemoveMovedKeysFromMetadata", True))
        skip = app_config.get("SkipColumns")
        if skip is None:
            skip = app_config.get("Data", {}).get("SkipColumns", [])
        self.skip = {s.lower() for s in skip}
        self._validate()
        self.key_by_lower = {k.lower(): k for k in self.columns}

    def _validate(self):
        errs = []
        if not _IDENT_RE.match(self.table):
            errs.append(f"TableName '{self.table}' must be letters/digits/_ and at most 63 chars")
        if not self.columns:
            errs.append("HorizontalMetadata.Columns is empty")
        if not 1 <= self.max_len <= 10485760:
            errs.append("ColumnMaxLength must be between 1 and 10485760")
        seen = {}
        for key, col in self.columns.items():
            if not key.strip():
                errs.append("HorizontalMetadata.Columns has an empty key")
            if not isinstance(col, str) or not _IDENT_RE.match(col):
                errs.append(f"Column name '{col}' (key {key}) must be letters/digits/_ and at most 63 chars")
            elif col in BASE_COLUMNS:
                errs.append(f"Column name '{col}' (key {key}) clashes with a built-in column")
            elif col.lower() in seen:
                errs.append(f"Column name '{col}' used by both {seen[col.lower()]} and {key}")
            else:
                seen[col.lower()] = key
        lower_keys = [k.lower() for k in self.columns]
        if len(set(lower_keys)) != len(lower_keys):
            errs.append("HorizontalMetadata.Columns has the same key twice (keys are case-insensitive)")
        if errs:
            raise ValueError("HorizontalMetadata config error:\n  " + "\n  ".join(errs))

    def horizontal_key(self, header):
        """Configured key for a source header, or None."""
        return self.key_by_lower.get(header.strip().lower())

    def is_skipped(self, header):
        return header.strip().lower() in self.skip


def q(ident):
    return '"' + ident.replace('"', '""') + '"'


# ---------------------------------------------------------------------------
# Table management
# ---------------------------------------------------------------------------

def ddl(cfg, with_sites=False):
    cols = ",\n".join(f"    {q(c)} varchar({cfg.max_len})" for c in cfg.columns.values())
    t = q(cfg.table)
    site_id = '    "SiteId" integer,\n' if with_sites else ""
    return f"""CREATE TABLE IF NOT EXISTS public.{t} (
    "MetadataHorizontalId" serial PRIMARY KEY,
    "FileId" integer NOT NULL UNIQUE REFERENCES public."Files"("FileId") ON DELETE CASCADE,
    "SiteName" varchar(256),
{site_id}{cols},
    "SourceFile" varchar(1000),
    "ImportedAt" timestamp NOT NULL DEFAULT CURRENT_TIMESTAMP,
    "UpdatedAt" timestamp NOT NULL DEFAULT CURRENT_TIMESTAMP
);"""


def _has_sites(conn):
    return conn.execute(text(
        "SELECT to_regclass('public.\"Sites\"') IS NOT NULL "
        "AND EXISTS (SELECT 1 FROM pg_proc WHERE proname = 'set_site_id')")).scalar()


def ensure_table(engine, cfg, log=print):
    """Create the table, or add columns that were added to the config since.
    Columns removed from the config are left in place (no data is dropped)."""
    with engine.begin() as conn:
        sites = _has_sites(conn)
        conn.execute(text(ddl(cfg, with_sites=sites)))
        existing = {r[0] for r in conn.execute(text(
            "SELECT column_name FROM information_schema.columns "
            "WHERE table_schema = 'public' AND table_name = :t"), {"t": cfg.table})}
        added = []
        for col in cfg.columns.values():
            if col not in existing:
                conn.execute(text(f"ALTER TABLE public.{q(cfg.table)} ADD COLUMN {q(col)} varchar({cfg.max_len})"))
                added.append(col)
        t = cfg.table
        conn.execute(text(f'CREATE INDEX IF NOT EXISTS {q("IX_" + t + "_SiteName")} ON public.{q(t)} ("SiteName")'))
        if sites:
            # same SiteId convention as the consolidated DB: FK to "Sites", kept by trigger
            conn.execute(text(f'ALTER TABLE public.{q(t)} ADD COLUMN IF NOT EXISTS "SiteId" integer'))
            conn.execute(text(f'CREATE INDEX IF NOT EXISTS {q("IX_" + t + "_SiteId")} ON public.{q(t)} ("SiteId")'))
            fk = t + "_SiteId_fkey"
            if not conn.execute(text("SELECT 1 FROM pg_constraint WHERE conname = :c"), {"c": fk}).first():
                conn.execute(text(f'ALTER TABLE public.{q(t)} ADD CONSTRAINT {q(fk)} '
                                  f'FOREIGN KEY ("SiteId") REFERENCES public."Sites"("SiteId")'))
            trg = "trg_" + t + "_site_id"
            if not conn.execute(text("SELECT 1 FROM pg_trigger WHERE tgname = :n"), {"n": trg}).first():
                conn.execute(text(f'CREATE TRIGGER {q(trg)} BEFORE INSERT OR UPDATE OF "SiteName" ON public.{q(t)} '
                                  f'FOR EACH ROW EXECUTE FUNCTION public.set_site_id()'))
        conn.execute(text(f"COMMENT ON TABLE public.{q(t)} IS "
                          f"'One row per file: fixed metadata columns from config.json HorizontalMetadata.Columns. "
                          f"All other metadata keys stay in \"Metadata\".'"))
        for key, col in cfg.columns.items():
            conn.execute(text(f"COMMENT ON COLUMN public.{q(t)}.{q(col)} IS :c"), {"c": f"Metadata key: {key}"})
    unused = sorted(existing - set(BASE_COLUMNS) - set(cfg.columns.values()))
    log(f"Table {cfg.table} ready ({len(cfg.columns)} metadata columns"
        + (f", added: {', '.join(added)}" if added else "")
        + (", SiteId linked to Sites" if sites else "") + ")")
    if unused:
        log(f"Note: columns no longer in config (kept, not dropped): {', '.join(unused)}")
    return added


# ---------------------------------------------------------------------------
# Reading CSV / XLSX
# ---------------------------------------------------------------------------

def _cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return _ILLEGAL.sub("", str(v))


def read_rows(path, sheet=None):
    """-> (headers, row_iterator). Values are strings ('' for empty)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = [_cell(h).strip() for h in next(it)]

        def gen():
            try:
                for r in it:
                    vals = [_cell(v) for v in r]
                    if any(vals):
                        yield vals
            finally:
                wb.close()
        return headers, gen()
    if ext in (".csv", ".txt"):
        f = open(path, "r", encoding="utf-8-sig", newline="")
        reader = csv.reader(f)
        headers = [h.strip() for h in next(reader)]

        def gen():
            try:
                for r in reader:
                    if any(r):
                        yield [_ILLEGAL.sub("", v) for v in r]
            finally:
                f.close()
        return headers, gen()
    raise ValueError(f"Unsupported file type: {ext} (use .csv or .xlsx)")


# ---------------------------------------------------------------------------
# Upload
# ---------------------------------------------------------------------------

def classify_headers(headers, cfg):
    """-> (horizontal {idx: key}, vertical {idx: header}, skipped [header], ids {name: idx})"""
    horiz, vert, skipped, ids = {}, {}, [], {}
    lower_ids = {n.lower(): n for n in ID_COLUMNS}
    for i, h in enumerate(headers):
        if not h:
            continue
        if h.lower() in lower_ids:
            ids.setdefault(lower_ids[h.lower()], i)
        key = cfg.horizontal_key(h)
        if key:
            horiz[i] = key
        elif cfg.is_skipped(h) or h.lower() in lower_ids:
            skipped.append(h)
        else:
            vert[i] = h[:256]
    return horiz, vert, skipped, ids


_STAGE_DDL = """
CREATE TEMP TABLE mh_stage (
    row_no      integer,
    file_id_in  text,
    unique_id   text,
    file_ref    text,
    orig_path   text,
    h           jsonb,
    v           jsonb,
    file_id     integer,
    matched_by  text
) ON COMMIT DROP;
"""


def import_file(engine, cfg, path, sheet=None, log=print, progress=None, report_dir=None):
    """Upload one CSV/XLSX. Returns a stats dict."""
    t0 = time.time()
    headers, rows = read_rows(path, sheet)
    horiz, vert, skipped, ids = classify_headers(headers, cfg)
    if not any(n in ids for n in ("FileId", "UniqueId", "FileRef", "OriginalSharePointPath")):
        raise ValueError("The file needs at least one of these columns to find the file: "
                         "FileId, UniqueId, FileRef, OriginalSharePointPath")
    log(f"{os.path.basename(path)}: {len(headers)} columns -> {len(horiz)} horizontal, "
        f"{len(vert)} vertical, {len(skipped)} skipped/identity")
    missing = sorted(set(cfg.columns) - set(horiz.values()))
    if missing:
        log(f"Configured keys not in this file (left unchanged): {', '.join(missing)}")

    ensure_table(engine, cfg, log=lambda m: None)
    stats = dict(rows_read=0, folders_skipped=0, versions_skipped=0, staged=0,
                 matched=0, unmatched=0, horizontal_rows=0, vertical_rows=0,
                 vertical_deleted=0, duplicates=0)

    def g(r, name):
        i = ids.get(name)
        return r[i].strip() if i is not None and i < len(r) else ""

    raw = engine.raw_connection()
    try:
        cur = raw.cursor()
        cur.execute(_STAGE_DDL)
        state = {"buf": io.StringIO(), "n": 0}
        state["w"] = csv.writer(state["buf"])

        def flush():
            state["buf"].seek(0)
            # FORCE_NOT_NULL: empty identity fields stay '' (not NULL) so the match rules work
            cur.copy_expert("COPY mh_stage (row_no, file_id_in, unique_id, file_ref, orig_path, h, v) "
                            "FROM STDIN WITH (FORMAT csv, "
                            "FORCE_NOT_NULL (file_id_in, unique_id, file_ref, orig_path))", state["buf"])
            state["buf"] = io.StringIO()
            state["w"] = csv.writer(state["buf"])
            state["n"] = 0

        for r in rows:
            stats["rows_read"] += 1
            if g(r, "FSObjType") == "1":
                stats["folders_skipped"] += 1
                continue
            vl = g(r, "VersionLabel")
            if cfg.current_only and vl and vl.lower() != "current":
                if vl.lower() == "folder":
                    stats["folders_skipped"] += 1
                else:
                    stats["versions_skipped"] += 1
                continue
            hv = {k: (r[i][:cfg.max_len] if i < len(r) else "") for i, k in horiz.items()}
            vv = {k: r[i][:4000] for i, k in vert.items() if i < len(r) and r[i] != ""}
            uid = g(r, "UniqueId").strip("{}").lower()
            state["w"].writerow([stats["rows_read"], g(r, "FileId"), uid, g(r, "FileRef"),
                                 g(r, "OriginalSharePointPath"),
                                 json.dumps(hv, ensure_ascii=False), json.dumps(vv, ensure_ascii=False)])
            state["n"] += 1
            stats["staged"] += 1
            if state["n"] >= STAGE_BATCH:
                flush()
                if progress:
                    progress(stats["rows_read"])
        if state["n"]:
            flush()
        log(f"Staged {stats['staged']:,} rows (read {stats['rows_read']:,}, folders {stats['folders_skipped']:,}, "
            f"old versions {stats['versions_skipped']:,})  [{time.time() - t0:.1f}s]")

        # ---- match rows to Files --------------------------------------------
        # FileId is trusted only when every other identity in the row agrees with
        # that file: FileIds were renumbered by the DB consolidation, so a FileId
        # from an older export may point at a different file.
        cur.execute("""
            UPDATE mh_stage s SET file_id = f."FileId", matched_by = 'FileId'
            FROM public."Files" f
            WHERE s.file_id_in ~ '^[0-9]{1,10}$' AND f."FileId" = s.file_id_in::bigint
              AND (s.unique_id = '' OR f."SharePointUniqueId" IS NULL
                   OR lower(trim(both '{}' from f."SharePointUniqueId")) = s.unique_id)
              AND (s.orig_path = '' OR f."OriginalSharePointPath" = s.orig_path)
              AND (s.file_ref = ''
                   OR regexp_replace(f."OriginalSharePointPath", '^[A-Za-z]+://[^/]+', '') = s.file_ref)""")
        cur.execute("""
            UPDATE mh_stage s SET file_id = f."FileId", matched_by = 'UniqueId'
            FROM public."Files" f
            WHERE s.file_id IS NULL AND s.unique_id <> ''
              AND f."SharePointUniqueId" = s.unique_id""")
        cur.execute("""
            UPDATE mh_stage s SET file_id = f."FileId", matched_by = 'UniqueId'
            FROM public."Files" f
            WHERE s.file_id IS NULL AND s.unique_id <> ''
              AND lower(trim(both '{}' from f."SharePointUniqueId")) = s.unique_id""")
        cur.execute("""
            UPDATE mh_stage s SET file_id = f."FileId", matched_by = 'OriginalSharePointPath'
            FROM public."Files" f
            WHERE s.file_id IS NULL AND s.orig_path <> ''
              AND f."OriginalSharePointPath" = s.orig_path""")
        cur.execute("""
            UPDATE mh_stage s SET file_id = f."FileId", matched_by = 'FileRef'
            FROM public."Files" f
            WHERE s.file_id IS NULL AND s.file_ref <> ''
              AND regexp_replace(f."OriginalSharePointPath", '^[A-Za-z]+://[^/]+', '') = s.file_ref""")
        cur.execute("SELECT matched_by, count(*) FROM mh_stage GROUP BY 1 ORDER BY 1 NULLS LAST")
        by = dict(cur.fetchall())
        stats["unmatched"] = by.pop(None, 0)
        stats["matched"] = sum(by.values())
        log("Matched " + ", ".join(f"{k}: {v:,}" for k, v in by.items())
            + f"; unmatched: {stats['unmatched']:,}")

        # duplicate rows for one file: the LAST row in the file wins
        cur.execute("""
            DELETE FROM mh_stage s USING mh_stage s2
            WHERE s.file_id = s2.file_id AND s.row_no < s2.row_no""")
        stats["duplicates"] = cur.rowcount
        if stats["duplicates"]:
            log(f"{stats['duplicates']:,} duplicate row(s) for the same file ignored (last row wins)")

        # ---- horizontal upsert ---------------------------------------------
        if horiz:
            present = [k for k in cfg.columns if k in horiz.values()]
            cols = [cfg.columns[k] for k in present]
            sel = ", ".join("NULLIF(s.h->>%s, '')" for _ in present)   # keys bound, never inlined
            upd = ", ".join(f"{q(c)} = EXCLUDED.{q(c)}" for c in cols)
            cur.execute(f"""
                INSERT INTO public.{q(cfg.table)} ("FileId", "SiteName", {", ".join(q(c) for c in cols)}, "SourceFile")
                SELECT s.file_id, f."SiteName", {sel}, %s
                FROM mh_stage s JOIN public."Files" f ON f."FileId" = s.file_id
                ON CONFLICT ("FileId") DO UPDATE SET {upd},
                    "SiteName" = EXCLUDED."SiteName", "SourceFile" = EXCLUDED."SourceFile",
                    "UpdatedAt" = CURRENT_TIMESTAMP""", (*present, os.path.basename(path)[:1000]))
            stats["horizontal_rows"] = cur.rowcount

        # ---- vertical: replace keys present in this file, insert the rest ---
        # exact spellings: config keys, plus the file's own header text for them
        keys = sorted(set(horiz.values()) | {headers[i] for i in horiz} | set(vert.values()))
        if keys:
            cur.execute("""
                DELETE FROM public."Metadata" m
                USING mh_stage s
                WHERE m."FileId" = s.file_id AND m."MetadataKey" = ANY(%s)""", (keys,))
            stats["vertical_deleted"] = cur.rowcount
        if vert:
            cur.execute("""
                INSERT INTO public."Metadata" ("FileId", "MetadataKey", "MetadataValue", "DataType", "SiteName")
                SELECT s.file_id, e.key, e.value, 'String', f."SiteName"
                FROM mh_stage s
                JOIN public."Files" f ON f."FileId" = s.file_id
                CROSS JOIN LATERAL jsonb_each_text(s.v) e
                WHERE e.value <> ''""")
            stats["vertical_rows"] = cur.rowcount

        # ---- unmatched report ------------------------------------------------
        if stats["unmatched"]:
            report_dir = report_dir or os.path.join(BASE_DIR, "logs")
            os.makedirs(report_dir, exist_ok=True)
            rpt = os.path.join(report_dir, f"unmatched_{os.path.splitext(os.path.basename(path))[0]}_"
                                            f"{datetime.now():%Y%m%d_%H%M%S}.csv")
            cur.execute("SELECT row_no, file_id_in, unique_id, file_ref, orig_path FROM mh_stage "
                        "WHERE file_id IS NULL ORDER BY row_no")
            with open(rpt, "w", encoding="utf-8-sig", newline="") as f:
                wr = csv.writer(f)
                wr.writerow(["SourceRowNo", "FileId", "UniqueId", "FileRef", "OriginalSharePointPath"])
                wr.writerows(cur.fetchall())
            stats["unmatched_report"] = rpt
            log(f"Unmatched rows written to {rpt}")

        raw.commit()
    except Exception:
        raw.rollback()
        raise
    finally:
        raw.close()

    stats["seconds"] = round(time.time() - t0, 1)
    log(f"Done: {stats['horizontal_rows']:,} horizontal rows upserted, {stats['vertical_rows']:,} vertical "
        f"Metadata rows written ({stats['vertical_deleted']:,} old rows replaced)  [{stats['seconds']}s]")
    return stats


# ---------------------------------------------------------------------------
# Convert existing vertical Metadata
# ---------------------------------------------------------------------------

def convert_existing(engine, cfg, remove_from_vertical=None, log=print):
    """Pivot configured keys already in "Metadata" into the horizontal table.
    Existing horizontal values are only overwritten by non-empty vertical values."""
    if remove_from_vertical is None:
        remove_from_vertical = cfg.remove_moved
    ensure_table(engine, cfg, log=log)
    t0 = time.time()
    keys = list(cfg.columns)
    cols = [cfg.columns[k] for k in keys]
    params = {f"k{i}": k for i, k in enumerate(keys)}      # keys bound, never inlined
    # exact key match so IX_Metadata_MetadataKey is used on large Metadata tables
    sel = ", ".join(f"MAX(CASE WHEN m.\"MetadataKey\" = :k{i} THEN m.\"MetadataValue\" END)"
                    for i in range(len(keys)))
    upd = ", ".join(f"{q(c)} = COALESCE(EXCLUDED.{q(c)}, t.{q(c)})" for c in cols)
    with engine.begin() as conn:
        res = conn.execute(text(f"""
            INSERT INTO public.{q(cfg.table)} AS t ("FileId", "SiteName", {", ".join(q(c) for c in cols)}, "SourceFile")
            SELECT m."FileId", MAX(f."SiteName"), {sel}, 'Metadata (converted)'
            FROM public."Metadata" m JOIN public."Files" f ON f."FileId" = m."FileId"
            WHERE m."MetadataKey" = ANY(:keys)
            GROUP BY m."FileId"
            ON CONFLICT ("FileId") DO UPDATE SET {upd}, "UpdatedAt" = CURRENT_TIMESTAMP"""),
            {"keys": keys, **params})
        moved = res.rowcount
        removed = 0
        if remove_from_vertical:
            removed = conn.execute(text(
                'DELETE FROM public."Metadata" WHERE "MetadataKey" = ANY(:keys)'),
                {"keys": keys}).rowcount
    log(f"Converted {moved:,} files into {cfg.table}"
        + (f"; removed {removed:,} rows from Metadata" if remove_from_vertical else "; Metadata left unchanged")
        + f"  [{time.time() - t0:.1f}s]")
    return {"files": moved, "removed": removed}


# ---------------------------------------------------------------------------
# Matrix: every distinct MetadataKey as a column
# ---------------------------------------------------------------------------

def matrix_table(app_config):
    t = (app_config.get("MetadataMatrix") or {}).get("TableName") or mm.DEFAULT_TABLE
    return mm.check_table(t, protected=(HConfig(app_config).table,))


def build_matrix(engine, app_config, log=print):
    """Rebuild the matrix table from "Metadata" (see metadata_matrix.py). One
    transaction, so the old matrix stays readable until the new one is complete."""
    t0 = time.time()
    table = matrix_table(app_config)
    with engine.begin() as conn:
        # plain DBAPI cursor, no parameters: the SQL holds keys as literals, which may
        # contain ':' or '%' (SQLAlchemy / psycopg2 would read those as placeholders)
        cur = conn.connection.cursor()

        def one(stmt):
            cur.execute(stmt)
            return cur.fetchone()[0]
        info = one(mm.info_sql(table))
        cols = mm.check_info(table, info)
        for line in mm.describe(info, cols):
            log(line)
        for stmt in mm.build_statements(table, info):
            cur.execute(stmt)
        summary = one(mm.summary_sql(table))
    log(mm.result_line(table, len(cols), summary, time.time() - t0))
    return {"files": summary["files"], "columns": len(cols), "duplicates": info["dups"],
            "orphan_files": summary["orphans"]}


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None):
    import argparse
    p = argparse.ArgumentParser(description="Horizontal metadata table tool")
    p.add_argument("command", choices=["ensure", "upload", "convert", "ddl", "matrix"])
    p.add_argument("file", nargs="?", help="CSV/XLSX for upload")
    p.add_argument("--sheet")
    p.add_argument("--keep-vertical", action="store_true", help="convert: do not delete moved keys from Metadata")
    p.add_argument("--config", default=CONFIG_FILE)
    p.add_argument("--url", default=os.environ.get("PG_URL"), help="default: config.json Connection + PGPASSWORD")
    a = p.parse_args(argv)
    app_config = load_config(a.config)
    cfg = HConfig(app_config)
    if a.command == "ddl":
        print(ddl(cfg))
        return 0
    engine = create_engine(a.url or connection_url(app_config))
    if a.command == "ensure":
        ensure_table(engine, cfg)
    elif a.command == "upload":
        if not a.file:
            p.error("upload needs a file")
        import_file(engine, cfg, a.file, sheet=a.sheet)
    elif a.command == "convert":
        convert_existing(engine, cfg, remove_from_vertical=not a.keep_vertical)
    elif a.command == "matrix":
        build_matrix(engine, app_config)
    return 0


if __name__ == "__main__":
    sys.exit(main())
